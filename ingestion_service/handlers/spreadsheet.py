"""Low-cost CSV and XLSX handlers."""

from __future__ import annotations

import csv
from io import StringIO
from tempfile import NamedTemporaryFile

from ingestion_service.chunking import chunk_rows
from ingestion_service.errors import ParseError
from ingestion_service.schemas import (
    DocumentHandlingPolicy,
    HandlerOutput,
    IngestedDocument,
    IngestedSection,
    SourceBlob,
)


class CsvHandler:
    name = "csv"
    supported_content_types = ("text/csv", "text/tab-separated-values")
    supported_extensions = (".csv", ".tsv")
    resource_tier = "minimal"

    async def parse(
        self,
        source: SourceBlob,
        *,
        policy: DocumentHandlingPolicy,
    ) -> HandlerOutput:
        delimiter = "\t" if source.source_uri.endswith(".tsv") else ","
        rows = [
            [cell.strip() for cell in row]
            for _, row in zip(
                range(policy.budget.max_rows),
                csv.reader(StringIO(source.text()), delimiter=delimiter),
            )
        ]
        return _spreadsheet_output(self.name, source, rows)


class XlsxHandler:
    name = "xlsx"
    supported_content_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroenabled.12",
    )
    supported_extensions = (".xlsx", ".xlsm")
    resource_tier = "minimal"

    async def parse(
        self,
        source: SourceBlob,
        *,
        policy: DocumentHandlingPolicy,
    ) -> HandlerOutput:
        try:
            from openpyxl import load_workbook  # type: ignore[import-not-found]
        except ModuleNotFoundError as exc:
            raise ParseError("XLSX parsing requires openpyxl") from exc
        rows: list[list[str]] = []
        with NamedTemporaryFile(suffix=".xlsx") as temp:
            temp.write(source.content)
            temp.flush()
            workbook = load_workbook(temp.name, read_only=True, data_only=True)
            for sheet in workbook.worksheets[: policy.budget.max_sheets]:
                for row in sheet.iter_rows(values_only=True):
                    rows.append(["" if cell is None else str(cell) for cell in row])
                    if len(rows) >= policy.budget.max_rows:
                        break
                if len(rows) >= policy.budget.max_rows:
                    break
            workbook.close()
        return _spreadsheet_output(self.name, source, rows)


def _spreadsheet_output(
    handler_name: str,
    source: SourceBlob,
    rows: list[list[str]],
) -> HandlerOutput:
    if not rows:
        raise ParseError("spreadsheet produced no rows")
    sections = [
        IngestedSection(
            section_id=str(index),
            text=text,
            order=index,
            metadata={"table_window": index},
        )
        for index, text in enumerate(chunk_rows(rows))
    ]
    document = IngestedDocument(
        document_id=source.document_id,
        source_uri=source.source_uri,
        content_type=source.content_type,
        data_type=source.data_type,
        content_hash=source.checksum,
        handler_name=handler_name,
        metadata=dict(source.metadata),
    )
    return HandlerOutput(document=document, sections=tuple(sections))
