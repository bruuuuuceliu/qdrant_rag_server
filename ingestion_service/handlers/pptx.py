"""Low-cost PPTX handler."""

from __future__ import annotations

from tempfile import NamedTemporaryFile

from ingestion_service.errors import ParseError
from ingestion_service.normalization import normalize_text
from ingestion_service.schemas import (
    DocumentHandlingPolicy,
    HandlerOutput,
    IngestedDocument,
    IngestedSection,
    SourceBlob,
)


class PptxNativeHandler:
    name = "pptx_native"
    supported_content_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    supported_extensions = (".pptx",)
    resource_tier = "minimal"

    async def parse(
        self,
        source: SourceBlob,
        *,
        policy: DocumentHandlingPolicy,
    ) -> HandlerOutput:
        try:
            from pptx import Presentation  # type: ignore[import-not-found]
        except ModuleNotFoundError as exc:
            raise ParseError("PPTX parsing requires python-pptx") from exc
        with NamedTemporaryFile(suffix=".pptx") as temp:
            temp.write(source.content)
            temp.flush()
            presentation = Presentation(temp.name)
        sections: list[IngestedSection] = []
        for index, slide in enumerate(presentation.slides[: policy.budget.max_slides]):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            text = normalize_text("\n".join(parts))
            if text:
                sections.append(
                    IngestedSection(
                        section_id=str(index),
                        text=text,
                        order=index,
                        metadata={"slide_number": index + 1},
                    )
                )
        if not sections:
            raise ParseError("PPTX produced no text")
        document = IngestedDocument(
            document_id=source.document_id,
            source_uri=source.source_uri,
            content_type=source.content_type,
            data_type=source.data_type,
            content_hash=source.checksum,
            handler_name=self.name,
            metadata=dict(source.metadata),
        )
        return HandlerOutput(document=document, sections=tuple(sections))
