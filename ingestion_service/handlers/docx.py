"""Low-cost DOCX handler."""

from __future__ import annotations

from tempfile import NamedTemporaryFile

from ingestion_service.errors import ParseError
from ingestion_service.normalization import normalize_text
from ingestion_service.schemas import (
    DocumentHandlingPolicy,
    HandlerOutput,
    IngestedDocument,
    IngestedSection,
    SourceBlob,
)


class DocxNativeHandler:
    name = "docx_native"
    supported_content_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    supported_extensions = (".docx",)
    resource_tier = "minimal"

    async def parse(
        self,
        source: SourceBlob,
        *,
        policy: DocumentHandlingPolicy,
    ) -> HandlerOutput:
        del policy
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ModuleNotFoundError as exc:
            raise ParseError("DOCX parsing requires python-docx") from exc
        with NamedTemporaryFile(suffix=".docx") as temp:
            temp.write(source.content)
            temp.flush()
            doc = Document(temp.name)
        sections = _paragraph_sections(doc)
        table_section = _table_section(doc, order=len(sections))
        if table_section is not None:
            sections.append(table_section)
        text = normalize_text("\n\n".join(section.text for section in sections))
        if not text:
            raise ParseError("DOCX produced no text")
        document = IngestedDocument(
            document_id=source.document_id,
            source_uri=source.source_uri,
            content_type=source.content_type,
            data_type=source.data_type,
            content_hash=source.checksum,
            handler_name=self.name,
            metadata=dict(source.metadata),
        )
        return HandlerOutput(document=document, sections=tuple(sections))


def _paragraph_sections(doc: object) -> list[IngestedSection]:
    sections: list[IngestedSection] = []
    page_parts: list[str] = []
    page_number = 1
    order = 0
    has_page_markers = False

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        break_count = _page_break_count(paragraph)
        if break_count:
            has_page_markers = True
        if text:
            page_parts.append(text)
        for _ in range(break_count):
            if page_parts:
                sections.append(
                    IngestedSection(
                        section_id=f"page-{page_number}",
                        text="\n\n".join(page_parts),
                        order=order,
                        page_number=page_number,
                    )
                )
                order += 1
                page_parts = []
            page_number += 1

    if page_parts:
        section_id = f"page-{page_number}" if has_page_markers else "paragraphs"
        sections.append(
            IngestedSection(
                section_id=section_id,
                text="\n\n".join(page_parts),
                order=order,
                page_number=page_number if has_page_markers else None,
            )
        )
    return sections


def _table_section(doc: object, *, order: int) -> IngestedSection | None:
    table_parts: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            table_parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    if not table_parts:
        return None
    return IngestedSection(
        section_id="tables",
        text="\n".join(table_parts),
        order=order,
    )


def _page_break_count(paragraph: object) -> int:
    from docx.oxml.ns import qn  # type: ignore[import-not-found]

    count = 0
    for br in paragraph._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            count += 1
    count += sum(1 for _ in paragraph._p.iter(qn("w:lastRenderedPageBreak")))
    return count
